from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "documentation" / "HireLens_Agents_Presentation_Final.pptx"
IMAGE = ROOT / "src" / "app" / "HireLens.jpg"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
FONT = "Segoe UI"
C = {
    "bg": "FFFFFF", "ink": "252525", "muted": "737373", "soft": "F7F7F7",
    "border": "E7E7E7", "dark": "202020", "white": "FFFFFF",
    "purple": "6657D9", "purple_soft": "F0EDFF", "green": "15803D",
    "green_soft": "ECFDF3", "amber": "B45309", "amber_soft": "FFF7E8",
    "red": "BE123C", "red_soft": "FFF1F2", "teal": "0F766E",
    "teal_soft": "ECFDFB", "blue": "1D4ED8", "blue_soft": "EFF6FF",
}


def rgb(name):
    return RGBColor.from_string(C[name])


def set_bg(slide, color="bg"):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def rect(slide, x, y, w, h, fill="soft", line="border", radius=True, width=1):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(width)
    if radius:
        try:
            shape.adjustments[0] = 0.12
        except Exception:
            pass
    return shape


def line(slide, x1, y1, x2, y2, color="border", width=1.2):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def text(slide, value, x, y, w, h, size=16, color="ink", bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def badge(slide, label, x, y, w=None, fill="soft", color="ink", border="border"):
    if w is None:
        w = max(0.85, len(label) * 0.075 + 0.35)
    rect(slide, x, y, w, 0.32, fill, border, True, 0.7)
    text(slide, label.upper(), x, y + 0.005, w, 0.29, 8.5, color, True,
         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def bullets(slide, items, x, y, w, h, size=13, gap=6):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = 0
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(gap)
        paragraph.line_spacing = 1.05
        bullet_run = paragraph.add_run()
        bullet_run.text = "• "
        bullet_run.font.name = FONT
        bullet_run.font.size = Pt(size)
        bullet_run.font.bold = True
        bullet_run.font.color.rgb = rgb("ink")
        head, body = item
        head_run = paragraph.add_run()
        head_run.text = head
        head_run.font.name = FONT
        head_run.font.size = Pt(size)
        head_run.font.bold = True
        head_run.font.color.rgb = rgb("ink")
        body_run = paragraph.add_run()
        body_run.text = body
        body_run.font.name = FONT
        body_run.font.size = Pt(size)
        body_run.font.color.rgb = rgb("muted")
    return box


def footer(slide, number, section="AGENT ARCHITECTURE"):
    line(slide, 0.55, 7.12, 12.78, 7.12, "border", 0.8)
    text(slide, "HireLens", 0.58, 7.17, 1.2, 0.18, 8.5, "muted", True)
    text(slide, section, 5.2, 7.17, 3.0, 0.18, 7.5, "muted", True, PP_ALIGN.CENTER)
    text(slide, f"{number:02d}", 12.25, 7.17, 0.5, 0.18, 8.5, "muted", True, PP_ALIGN.RIGHT)


def heading(slide, kicker, title, subtitle, number):
    set_bg(slide)
    badge(slide, kicker, 0.6, 0.38)
    text(slide, title, 0.6, 0.87, 12.0, 0.58, 28, "ink", True)
    text(slide, subtitle, 0.6, 1.48, 11.9, 0.48, 12.5, "muted")


def node(slide, x, y, w, title, fill="soft", color="ink"):
    rect(slide, x, y, w, 0.74, fill, "border", True, 0.8)
    text(slide, title, x + 0.1, y + 0.22, w - 0.2, 0.25, 10.5, color, True, PP_ALIGN.CENTER)


def arrow(slide, x1, y, x2):
    shape = line(slide, x1, y, x2, y, "muted", 1.4)
    shape.line.end_arrowhead = True


def agent_slide(number, name, role, accent, consumes, work, produces, controls, flow):
    slide = prs.slides.add_slide(BLANK)
    heading(slide, "SPECIALISED AGENT", name, role, number)
    rect(slide, 0.6, 2.1, 1.05, 4.62, "dark", "dark")
    text(slide, f"{number - 4:02d}", 0.6, 2.26, 1.05, 0.55, 30, "white", True, PP_ALIGN.CENTER)
    text(slide, "AGENT", 0.6, 2.86, 1.05, 0.25, 9, "white", True, PP_ALIGN.CENTER)
    line(slide, 0.82, 3.28, 1.43, 3.28, "muted", 1)
    text(slide, flow, 0.73, 3.48, 0.8, 2.7, 9, "white", False, PP_ALIGN.CENTER)
    widths = [2.62, 3.9, 2.62]
    xs = [1.88, 4.72, 8.84]
    labels = ["CONSUMES", "CORE WORK", "PRODUCES"]
    for x, width, label, items in zip(xs, widths, labels, [consumes, work, produces]):
        rect(slide, x, 2.1, width, 3.15)
        badge(slide, label, x + 0.2, 2.28, fill=f"{accent}_soft", color=accent, border=f"{accent}_soft")
        bullets(slide, items, x + 0.2, 2.78, width - 0.4, 2.2, 12.3, 5)
    rect(slide, 1.88, 5.48, 9.58, 1.24, "bg", "border")
    badge(slide, "GUARDRAILS & ORCHESTRATION", 2.08, 5.67, 2.05)
    bullets(slide, controls, 2.12, 6.02, 9.0, 0.52, 11.3, 2)
    footer(slide, number)


# 1 — Title
slide = prs.slides.add_slide(BLANK)
set_bg(slide)
rect(slide, 0.6, 0.55, 6.7, 6.35)
badge(slide, "AGENTIC RECRUITMENT INTELLIGENCE", 0.95, 0.95, 2.65, "dark", "white", "dark")
text(slide, "HireLens", 0.95, 1.55, 5.7, 0.7, 34, "ink", True)
text(slide, "The multi-agent system\nbehind evidence-led hiring", 0.95, 2.32, 5.8, 1.2, 26, "ink", True)
text(slide, "Eight specialised agents. Four LangGraph workflows.\nOne human decision boundary.", 0.95, 3.72, 5.65, 0.68, 15, "muted")
line(slide, 0.95, 4.62, 6.65, 4.62)
for i, (label, value) in enumerate([("AGENTS", "08"), ("WORKFLOWS", "04"), ("DECISION OWNER", "HUMAN")]):
    x = 0.95 + i * 1.88
    text(slide, value, x, 4.9, 1.65, 0.42, 18, "ink", True)
    text(slide, label, x, 5.36, 1.65, 0.25, 8, "muted", True)
text(slide, "Architecture overview · 2026", 0.95, 6.25, 4, 0.25, 10, "muted")
if IMAGE.exists():
    slide.shapes.add_picture(str(IMAGE), Inches(7.7), Inches(1.05), Inches(5), Inches(5))
text(slide, "Privacy-aware  •  Observable  •  Advisory", 7.8, 6.25, 4.8, 0.3, 11, "muted", True, PP_ALIGN.CENTER)
footer(slide, 1, "HIRELENS OVERVIEW")

# 2 — Purpose
slide = prs.slides.add_slide(BLANK)
heading(slide, "SYSTEM PURPOSE", "From documents to defensible decisions", "HireLens turns hiring inputs into structured, traceable evidence while keeping the final decision human.", 2)
principles = [
    ("01", "STRUCTURE", "Convert JDs, resumes and transcripts into validated domain records.", "blue"),
    ("02", "GROUND", "Retrieve comparable, expert-verified cases before model reasoning.", "purple"),
    ("03", "VALIDATE", "Mechanically check quotes, case IDs, schemas and severity consistency.", "teal"),
    ("04", "HAND OFF", "Package recommendations, uncertainty and open questions for a reviewer.", "amber"),
]
for i, (label, title, body, accent) in enumerate(principles):
    x = 0.6 + i * 3.13
    rect(slide, x, 2.22, 2.85, 3.23)
    badge(slide, label, x + 0.22, 2.45, 0.48, f"{accent}_soft", accent, f"{accent}_soft")
    text(slide, title, x + 0.22, 3.02, 2.4, 0.4, 17, "ink", True)
    text(slide, body, x + 0.22, 3.62, 2.38, 1.14, 13, "muted")
rect(slide, 0.6, 5.75, 12.13, 0.9, "dark", "dark")
text(slide, "CORE PRINCIPLE", 0.86, 5.98, 1.45, 0.22, 9, "white", True)
text(slide, "Agents advise. Evidence constrains. Humans decide.", 2.25, 5.88, 9.85, 0.36, 19, "white", True)
footer(slide, 2, "SYSTEM PURPOSE")

# 3 — Workflow map
slide = prs.slides.add_slide(BLANK)
heading(slide, "ORCHESTRATION", "Four workflows, one evidence memory", "A deterministic control plane routes state through specialised agents and durable handoffs.", 3)
flows = [
    ("JD INTAKE", [("Load", .9), ("JD Agent", 1.35), ("Persist", 1.05), ("Index", .95)], "blue"),
    ("RESUME INTAKE", [("Load", .9), ("Resume Agent", 1.5), ("Persist", 1.05), ("Index", .95)], "teal"),
    ("CANDIDATE ANALYSIS", [("Supervisor", 1.25), ("RAG plan", 1.15), ("Gap Agent", 1.25), ("RAG gaps", 1.15), ("Question", 1.2)], "purple"),
    ("TRANSCRIPT REVIEW", [("Eval RAG", 1.15), ("Transcript", 1.25), ("Flag RAG", 1.15), ("Red Flag", 1.15), ("Review", 1.1)], "amber"),
]
for y, (label, nodes, accent) in zip([2.14, 3.25, 4.36, 5.47], flows):
    badge(slide, label, 0.62, y + 0.2, 1.55, f"{accent}_soft", accent, f"{accent}_soft")
    x = 2.4
    for index, (name, width) in enumerate(nodes):
        node(slide, x, y, width, name, f"{accent}_soft", accent)
        if index < len(nodes) - 1:
            arrow(slide, x + width + 0.06, y + 0.37, x + width + 0.30)
        x += width + 0.42
text(slide, "Shared services", 10.72, 2.15, 1.5, 0.25, 10, "muted", True)
for index, label in enumerate(["pgvector memory", "Reflexion", "Agent memory", "Human feedback"]):
    badge(slide, label, 10.7, 2.55 + index * 0.68, 1.55)
footer(slide, 3, "WORKFLOW MAP")

# 4 — Roster
slide = prs.slides.add_slide(BLANK)
heading(slide, "AGENT ROSTER", "Eight agents, deliberately narrow roles", "Specialisation keeps extraction, retrieval, reasoning, validation and review independently observable.", 4)
agents = [
    ("01", "JD Agent", "Requirements extraction", "blue"), ("02", "Resume Agent", "Privacy-filtered profile", "teal"),
    ("03", "Evidence Retrieval", "Intent + semantic context", "purple"), ("04", "Gap Analysis", "Fit reasoning + gaps", "amber"),
    ("05", "Question Agent", "Targeted interview plan", "blue"), ("06", "Transcript Evaluation", "Answer scoring + quotes", "teal"),
    ("07", "Red Flag Agent", "Case-grounded concerns", "red"), ("08", "Human Review Agent", "Decision packet", "purple"),
]
for index, (number, name, description, accent) in enumerate(agents):
    x = 0.6 + (index % 4) * 3.08
    y = 2.12 + (index // 4) * 2.08
    rect(slide, x, y, 2.82, 1.78)
    badge(slide, number, x + 0.2, y + 0.2, 0.48, f"{accent}_soft", accent, f"{accent}_soft")
    text(slide, name, x + 0.2, y + 0.66, 2.4, 0.35, 15, "ink", True)
    text(slide, description, x + 0.2, y + 1.13, 2.4, 0.32, 11.5, "muted")
footer(slide, 4, "AGENT ROSTER")

# 5–12 — Individual agents
agent_slide(5, "JD Agent", "Turns an unstructured role description into a weighted requirement contract.", "blue",
    [("Input: ", "raw job description"), ("Context: ", "title, seniority, domain")],
    [("Extracts ", "skills, technologies, experience, certifications and domains"), ("Weights ", "importance from 0–100"), ("Classifies ", "required vs. nice-to-have")],
    [("Structured JD ", "validated by Zod"), ("Requirement rows ", "for deterministic matching"), ("Privacy-safe vectors ", "for retrieval")],
    [("Structured output — ", "prevents free-form drift"), ("Persistence + indexing — ", "occur after validation")],
    "JD intake\nload → extract\n→ persist → index")

agent_slide(6, "Resume Agent", "Builds an anonymous, evidence-bearing candidate profile.", "teal",
    [("Input: ", "resume text"), ("Local fields: ", "name, email, phone")],
    [("Redacts ", "identity, contact, URLs, location and institutions"), ("Extracts ", "skills with proficiency and evidence"), ("Normalises ", "employment, projects and total experience")],
    [("Candidate profile ", "without outbound PII"), ("Skill evidence ", "for matching"), ("Anonymous vectors ", "for future retrieval")],
    [("Original source — ", "retained for authorised users"), ("Demo preview — ", "masked before display"), ("AI boundary — ", "always centrally sanitised")],
    "Resume intake\nlocal contact → redact\n→ extract → persist")

agent_slide(7, "Evidence Retrieval Agent", "Resolves intent before retrieving semantic context for reasoning.", "purple",
    [("Purpose: ", "candidate fit, gap validation, interview evaluation or red flags"), ("Sources: ", "knowledge entries + evaluations")],
    [("Plans ", "2–4 privacy-safe semantic queries"), ("Searches ", "pgvector for comparable outcomes"), ("Deduplicates ", "by case owner and ranks similarity"), ("Returns ", "stable case IDs, excerpts and outcomes")],
    [("Grounding packet ", "injected into reasoning prompts"), ("Retrieval trace ", "intent + winning query"), ("Fallback plan ", "if intent planning fails")],
    [("No hiring answer — ", "during planning"), ("Current candidate — ", "excluded from own precedent"), ("PII — ", "redacted before planning and embedding")],
    "Agentic RAG\nintent → multi-query\n→ rank → reason")

agent_slide(8, "Gap Analysis Agent", "Combines deterministic coverage with qualitative fit reasoning.", "amber",
    [("Job: ", "weighted requirements"), ("Candidate: ", "skills + evidence"), ("Context: ", "peer summary, RAG cases, calibration")],
    [("Compares ", "strengths, partial matches and missing skills"), ("Explains ", "match score and verdict"), ("Identifies ", "areas requiring interview validation"), ("Differentiates ", "against up to five pipeline peers")],
    [("Match analysis ", "score + verdict"), ("Gap set ", "severity and importance"), ("Validation areas ", "for the Question Agent")],
    [("Supervisor depth — ", "minimal / standard / deep"), ("Coverage anchor — ", "limits model drift"), ("Memory notes — ", "adjust recurring blind spots")],
    "Candidate analysis\ncoverage + evidence\n→ reason → gaps")

agent_slide(9, "Question Agent", "Converts verified gaps into an interview validation plan.", "blue",
    [("Input: ", "gap analysis"), ("Grounding: ", "targeted gap-evidence cases"), ("Profile: ", "candidate history + role needs")],
    [("Generates ", "screening, technical, gap and experience questions"), ("Targets ", "strong claims, uncertainty and missing evidence"), ("Defines ", "expected signals and rationale"), ("Calibrates ", "difficulty to role and candidate")],
    [("Question set ", "candidate-specific"), ("Expected signals ", "for later scoring"), ("Rationale ", "why each question matters")],
    [("No generic bank dump — ", "questions follow observed evidence"), ("Persisted set — ", "becomes reusable and retrievable")],
    "Candidate analysis\ngaps + precedent\n→ targeted questions")

agent_slide(10, "Transcript Evaluation Agent", "Scores demonstrated interview evidence against expected signals.", "teal",
    [("Input: ", "raw transcript + participants"), ("Plan: ", "question set + expected signals"), ("RAG: ", "interview-evaluation cases")],
    [("Scores ", "technical depth and communication separately"), ("Maps ", "signals hit and missed per question"), ("Quotes ", "the transcript for every judgement"), ("Retries ", "when validation detects unsupported quotes")],
    [("0–10 ratings ", "technical, communication, overall"), ("Answer breakdown ", "question by question"), ("Recommendation ", "with rationale and evidence")],
    [("Quote validator — ", "checks source text mechanically"), ("Reflexion — ", "maximum two attempts"), ("Durable artifact — ", "versioned and auditable")],
    "Transcript review\nquestions + transcript\n→ score + validate")

agent_slide(11, "Red Flag Agent", "Surfaces only concerns grounded in candidate evidence and precedent.", "red",
    [("Input: ", "resume claims + transcript"), ("RAG: ", "dedicated red-flag precedents"), ("Context: ", "role requirements + calibration")],
    [("Checks ", "seniority, project depth, contradictions, claims and timelines"), ("Requires ", "direct source quote for each flag"), ("Cites ", "retrieved case ID + relevance"), ("Frames ", "every concern as something to verify")],
    [("GREEN / YELLOW / RED ", "advisory severity"), ("Evidence bundle ", "source quotes + precedent IDs"), ("Follow-up concern ", "never an automatic rejection")],
    [("Mechanical validation — ", "quote exists in claimed source"), ("Case validation — ", "ID must be actually retrieved"), ("Reflexion — ", "retries invalid grounding")],
    "Transcript review\nflag RAG → detect\n→ ground → validate")

agent_slide(12, "Human Review Agent", "Synthesises evidence into a decision packet—not a decision.", "purple",
    [("Inputs: ", "match, transcript evaluation, flags"), ("Context: ", "comparable cases + reviewer memory")],
    [("Weighs ", "scores, concerns and uncertainty"), ("Summarises ", "key evidence and unresolved questions"), ("Explains ", "why precedent is relevant"), ("Routes ", "final judgement to a person")],
    [("Recommendation ", "advance / hold / reject"), ("Headline + reasoning ", "review-ready"), ("Open questions ", "what a human must resolve")],
    [("GREEN optimisation — ", "deterministic packet when no flags"), ("Human authority — ", "only reviewer changes status"), ("Overrides — ", "feed calibration memory")],
    "Review synthesis\nvalidated artifacts\n→ human packet")

# 13 — Agentic RAG
slide = prs.slides.add_slide(BLANK)
heading(slide, "DEEP DIVE", "Agentic RAG: plan, retrieve, reason, learn", "Retrieval is a first-class agent step—not a one-shot lookup attached to a prompt.", 13)
steps = [
    ("1", "RESOLVE INTENT", "Typed purpose: fit, gaps, interview or red flags", "purple"),
    ("2", "PLAN QUERIES", "Generate 2–4 diverse, privacy-safe searches", "blue"),
    ("3", "RETRIEVE", "Embed and search expert-verified historical cases", "teal"),
    ("4", "RANK + DEDUPE", "Keep strongest chunk per owner; preserve stable IDs", "amber"),
    ("5", "REASON", "Inject context into the specialised reasoning agent", "purple"),
    ("6", "LEARN", "Index reviewer decisions and verified gaps", "green"),
]
for index, (number, title, body, accent) in enumerate(steps):
    x = 0.6 + (index % 3) * 4.08
    y = 2.08 + (index // 3) * 2.05
    rect(slide, x, y, 3.8, 1.72)
    badge(slide, number, x + 0.2, y + 0.2, 0.45, f"{accent}_soft", accent, f"{accent}_soft")
    text(slide, title, x + 0.78, y + 0.2, 2.75, 0.28, 13, "ink", True)
    text(slide, body, x + 0.2, y + 0.72, 3.35, 0.62, 11.5, "muted")
footer(slide, 13, "AGENTIC RAG")

# 14 — Trust loops
slide = prs.slides.add_slide(BLANK)
heading(slide, "TRUST LAYER", "Evidence, reflexion and expert feedback", "Three loops tighten quality without allowing models to become the decision-maker.", 14)
loops = [
    ("MECHANICAL VALIDATION", "Quotes must exist in source text. Red-flag case IDs must appear in retrieval results. Severity must be internally consistent.", "red"),
    ("REFLEXION", "Validation errors return as explicit feedback. The agent retries once; repeated failure becomes a bias-warning memory note.", "amber"),
    ("EXPERT FEEDBACK", "Reviewers record decisions, overrides and verified gaps. They are embedded as future precedent and written to calibration memory.", "green"),
]
for index, (title, body, accent) in enumerate(loops):
    x = 0.6 + index * 4.08
    rect(slide, x, 2.2, 3.8, 3.65)
    badge(slide, f"LOOP {index + 1}", x + 0.22, 2.45, 0.78, f"{accent}_soft", accent, f"{accent}_soft")
    text(slide, title, x + 0.22, 3.0, 3.25, 0.38, 16, "ink", True)
    text(slide, body, x + 0.22, 3.58, 3.3, 1.5, 13, "muted")
rect(slide, 0.6, 6.08, 12.13, 0.62, "dark", "dark")
text(slide, "Result: evidence improves over time, while the human remains the only hiring authority.", 0.85, 6.23, 11.55, 0.25, 14, "white", True, PP_ALIGN.CENTER)
footer(slide, 14, "TRUST & LEARNING")

# 15 — Privacy and observability
slide = prs.slides.add_slide(BLANK)
heading(slide, "OPERATING MODEL", "Private by boundary. Observable by design.", "Demo mode explains what is running without weakening the production privacy boundary.", 15)
rect(slide, 0.6, 2.15, 5.85, 4.25)
badge(slide, "PRIVACY BOUNDARY", 0.85, 2.42, 1.45, "green_soft", "green", "green_soft")
bullets(slide, [
    ("Local preservation — ", "authorised users retain original resumes and transcripts."),
    ("Demo masking — ", "uploaded PII is redacted before textarea display."),
    ("Outbound sanitisation — ", "names, contacts, locations, institutions and URLs are removed."),
    ("Vector hygiene — ", "content and queries are privacy-filtered before embedding."),
], 0.85, 3.02, 5.25, 2.7, 13, 8)
rect(slide, 6.75, 2.15, 5.98, 4.25)
badge(slide, "UNDER THE HOOD", 7.0, 2.42, 1.35, "purple_soft", "purple", "purple_soft")
bullets(slide, [
    ("Running-agent history — ", "timestamped, persistent and manually reset."),
    ("Workflow path — ", "running, complete and failed lifecycle states."),
    ("Token visibility — ", "optional prompt, completion and total usage."),
    ("Prompt controls — ", "Demo overrides still pass central PII filtering."),
], 7.0, 3.02, 5.25, 2.7, 13, 8)
footer(slide, 15, "PRIVACY & OBSERVABILITY")

# 16 — Closing
slide = prs.slides.add_slide(BLANK)
set_bg(slide)
rect(slide, 0.6, 0.58, 12.13, 6.25, "dark", "dark")
badge(slide, "HIRELENS", 0.98, 0.98, 1.02, "white", "ink", "white")
text(slide, "Specialised agents.\nGrounded evidence.\nHuman decisions.", 0.98, 1.65, 7.2, 2.25, 32, "white", True)
text(slide, "A recruitment intelligence architecture designed to make\nAI useful, inspectable and appropriately constrained.", 0.98, 4.35, 6.9, 0.8, 15, "white")
if IMAGE.exists():
    slide.shapes.add_picture(str(IMAGE), Inches(8.45), Inches(1.35), Inches(3.55), Inches(3.55))
badge(slide, "PRIVACY-AWARE", 8.48, 5.28, 1.25, "green_soft", "green", "green_soft")
badge(slide, "EVIDENCE-LED", 9.9, 5.28, 1.18, "purple_soft", "purple", "purple_soft")
badge(slide, "HUMAN-OWNED", 11.23, 5.28, 1.12, "amber_soft", "amber", "amber_soft")
footer(slide, 16, "ARCHITECTURE SUMMARY")

# 17 — Business problem
slide = prs.slides.add_slide(BLANK)
heading(slide, "BUSINESS PROBLEM", "Hiring evidence is fragmented, variable and hard to audit", "The challenge is not a lack of data—it is converting inconsistent documents and interviews into comparable, defensible decisions.", 17)
problems = [
    ("FRAGMENTED INPUTS", "JDs, resumes, transcripts, notes and historical outcomes live in different formats and levels of detail.", "red"),
    ("INCONSISTENT REVIEW", "Different reviewers apply different standards, questions and interpretations to similar candidates.", "amber"),
    ("WEAK TRACEABILITY", "Scores and concerns are difficult to reconstruct when the supporting evidence is not captured with the decision.", "purple"),
    ("PRIVACY PRESSURE", "Recruitment data contains identity, contact and education details that should not cross every AI boundary.", "teal"),
]
for index, (title, body, accent) in enumerate(problems):
    x = 0.6 + (index % 2) * 6.15
    y = 2.12 + (index // 2) * 2.05
    rect(slide, x, y, 5.85, 1.76)
    badge(slide, f"0{index + 1}", x + 0.22, y + 0.22, 0.5, f"{accent}_soft", accent, f"{accent}_soft")
    text(slide, title, x + 0.9, y + 0.2, 4.55, 0.3, 14, "ink", True)
    text(slide, body, x + 0.22, y + 0.72, 5.35, 0.68, 12, "muted")
rect(slide, 0.6, 6.24, 12.13, 0.48, "dark", "dark")
text(slide, "Business requirement: faster synthesis without surrendering evidence quality, privacy or human accountability.", 0.85, 6.35, 11.62, 0.2, 12.5, "white", True, PP_ALIGN.CENTER)
footer(slide, 17, "BUSINESS CASE")

# 18 — Why agentic
slide = prs.slides.add_slide(BLANK)
heading(slide, "WHY AGENTIC", "Use automation, agents and humans where each is strongest", "HireLens is deliberately hybrid: rules provide repeatability, agents handle contextual reasoning, and people own consequential judgement.", 18)
columns = [
    ("DETERMINISTIC AUTOMATION", "Best for", [("Rules — ", "schema checks, scoring and persistence"), ("Strength — ", "speed and repeatability"), ("Limit — ", "cannot resolve ambiguous evidence or intent")], "blue"),
    ("AGENTIC AI", "Best for", [("Reasoning — ", "intent, retrieval, gaps and synthesis"), ("Strength — ", "adapts to unstructured context"), ("Limit — ", "must be grounded and validated")], "purple"),
    ("HUMAN EXPERT", "Best for", [("Judgement — ", "trade-offs, fairness and final decision"), ("Strength — ", "accountability and domain nuance"), ("Limit — ", "time, consistency and cognitive load")], "green"),
]
for index, (title, sub, items, accent) in enumerate(columns):
    x = 0.6 + index * 4.08
    rect(slide, x, 2.15, 3.8, 3.82)
    badge(slide, sub, x + 0.22, 2.4, 0.66, f"{accent}_soft", accent, f"{accent}_soft")
    text(slide, title, x + 0.22, 2.92, 3.3, 0.38, 15, "ink", True)
    bullets(slide, items, x + 0.22, 3.55, 3.3, 1.9, 12.5, 9)
rect(slide, 0.6, 6.2, 12.13, 0.54, "purple_soft", "purple_soft")
text(slide, "Agentic is justified where the system must decide what evidence to seek next—not merely execute a fixed lookup.", 0.85, 6.33, 11.6, 0.23, 13, "purple", True, PP_ALIGN.CENTER)
footer(slide, 18, "WHY AGENTIC")

# 19 — Technical architecture
slide = prs.slides.add_slide(BLANK)
heading(slide, "TECHNICAL ARCHITECTURE", "A layered, auditable agent platform", "Next.js coordinates user workflows; LangGraph manages state; Supabase persists evidence, vectors, tasks and feedback.", 19)
layers = [
    ("EXPERIENCE", ["Dashboard", "Jobs", "Candidates", "Review", "Demo observability"], "blue"),
    ("API + CONTROL", ["Authenticated routes", "Zod contracts", "Deterministic scoring", "PII boundary"], "teal"),
    ("LANGGRAPH", ["JD intake", "Resume intake", "Candidate analysis", "Transcript review"], "purple"),
    ("AGENT SERVICES", ["Structured LLMs", "Intent planner", "Reflexion", "Validation"], "amber"),
    ("DATA + MEMORY", ["Postgres", "pgvector", "Agent memory", "Tasks / artifacts / events"], "green"),
]
for index, (label, items, accent) in enumerate(layers):
    y = 2.05 + index * 0.87
    badge(slide, label, 0.62, y + 0.18, 1.42, f"{accent}_soft", accent, f"{accent}_soft")
    rect(slide, 2.28, y, 10.45, 0.68, "soft", "border")
    item_w = 9.85 / len(items)
    for j, item in enumerate(items):
        if j:
            line(slide, 2.58 + j * item_w, y + 0.13, 2.58 + j * item_w, y + 0.55, "border", 0.8)
        text(slide, item, 2.52 + j * item_w, y + 0.21, item_w - 0.08, 0.25, 10.5, "ink", True, PP_ALIGN.CENTER)
footer(slide, 19, "TECHNICAL ARCHITECTURE")

# 20 — Evaluation and edge cases
slide = prs.slides.add_slide(BLANK)
heading(slide, "EVALUATION", "Test outputs, grounding and failure modes", "Evaluation must cover both model quality and orchestration correctness—not only happy-path accuracy.", 20)
rect(slide, 0.6, 2.12, 5.85, 4.35)
badge(slide, "AUTOMATED EVALUATION", 0.85, 2.4, 1.55, "purple_soft", "purple", "purple_soft")
bullets(slide, [
    ("Schema tests — ", "all structured outputs satisfy strict Zod contracts."),
    ("Grounding tests — ", "quotes exist in source text and case IDs were retrieved."),
    ("Workflow tests — ", "branch joins, idempotency and lifecycle state remain correct."),
    ("Offline quality — ", "DeepEval hallucination checks use exact supplied context."),
], 0.85, 3.02, 5.25, 2.85, 12.4, 8)
rect(slide, 6.75, 2.12, 5.98, 4.35)
badge(slide, "EDGE-CASE SUITE", 7.0, 2.4, 1.22, "red_soft", "red", "red_soft")
bullets(slide, [
    ("Cold retrieval — ", "no vectors or low-similarity cases."),
    ("Noisy documents — ", "OCR artefacts, missing sections and malformed VTT."),
    ("Conflicting claims — ", "resume/transcript dates, roles and technologies disagree."),
    ("Adversarial grounding — ", "invented quotes, case IDs, PII and prompt injection."),
    ("Operational faults — ", "timeouts, retries, duplicate requests and partial branch failure."),
], 7.0, 3.02, 5.25, 3.0, 12.2, 7)
footer(slide, 20, "EVALUATION")

# 21 — HITL validation
slide = prs.slides.add_slide(BLANK)
heading(slide, "HUMAN-IN-THE-LOOP", "Expert validation closes the learning loop", "The reviewer does more than approve an answer: their correction becomes evidence for the next retrieval and calibration cycle.", 21)
steps = [
    ("1", "REVIEW", "Inspect scores, quotes, flags and precedent.", "blue"),
    ("2", "DECIDE", "Accept or override; only a person changes status.", "purple"),
    ("3", "VERIFY GAPS", "Record explicit domain gaps, one per line.", "amber"),
    ("4", "INDEX", "Embed decision, reasoning and verified gaps.", "teal"),
    ("5", "CALIBRATE", "Write high-confidence Gap Analysis memory.", "green"),
    ("6", "RETRIEVE AGAIN", "Future agents recover expert precedent.", "purple"),
]
for index, (num, title, body, accent) in enumerate(steps):
    x = 0.6 + (index % 3) * 4.08
    y = 2.1 + (index // 3) * 2.05
    rect(slide, x, y, 3.8, 1.7)
    badge(slide, num, x + 0.2, y + 0.2, 0.44, f"{accent}_soft", accent, f"{accent}_soft")
    text(slide, title, x + 0.78, y + 0.2, 2.75, 0.28, 13, "ink", True)
    text(slide, body, x + 0.2, y + 0.72, 3.35, 0.58, 11.5, "muted")
rect(slide, 0.6, 6.24, 12.13, 0.48, "dark", "dark")
text(slide, "Governance rule: model recommendations are advisory; expert feedback is attributable, persistent and retrievable.", 0.85, 6.35, 11.62, 0.2, 12.5, "white", True, PP_ALIGN.CENTER)
footer(slide, 21, "HUMAN VALIDATION")

# 22 — ROI
slide = prs.slides.add_slide(BLANK)
heading(slide, "ROI FRAMEWORK", "Measure value without inventing it", "Establish a baseline, instrument the workflow, and report observed improvement by role, stage and candidate volume.", 22)
drivers = [
    ("TIME", "Minutes to structure a role, screen evidence, prepare questions and assemble review packets.", "blue"),
    ("QUALITY", "Quote-grounding rate, retrieval relevance, override rate and unresolved evidence gaps.", "purple"),
    ("CONSISTENCY", "Variance in scoring, question coverage and decisions across reviewers and jobs.", "teal"),
    ("RISK", "PII boundary violations, unsupported claims, audit gaps and rework after incorrect flags.", "red"),
]
for index, (title, body, accent) in enumerate(drivers):
    x = 0.6 + index * 3.08
    rect(slide, x, 2.15, 2.82, 2.35)
    badge(slide, f"DRIVER {index + 1}", x + 0.2, 2.4, 0.82, f"{accent}_soft", accent, f"{accent}_soft")
    text(slide, title, x + 0.2, 2.95, 2.38, 0.34, 16, "ink", True)
    text(slide, body, x + 0.2, 3.48, 2.38, 0.72, 11.5, "muted")
rect(slide, 0.6, 4.82, 7.35, 1.62, "soft", "border")
badge(slide, "MEASUREMENT MODEL", 0.84, 5.05, 1.3, "green_soft", "green", "green_soft")
text(slide, "ROI = (labour saved + avoided rework + quality/risk value − operating cost) ÷ operating cost", 0.84, 5.55, 6.85, 0.48, 14, "ink", True, PP_ALIGN.CENTER)
rect(slide, 8.2, 4.82, 4.53, 1.62, "dark", "dark")
text(slide, "REPORT ONLY OBSERVED RESULTS", 8.48, 5.08, 3.95, 0.25, 10, "white", True, PP_ALIGN.CENTER)
text(slide, "Baseline → pilot cohort → production trend", 8.48, 5.52, 3.95, 0.32, 15, "white", True, PP_ALIGN.CENTER)
footer(slide, 22, "ROI")

prs.core_properties.title = "HireLens — Agentic AI Agents"
prs.core_properties.subject = "Application agent architecture, Agentic RAG, privacy and observability"
prs.core_properties.author = "HireLens"
prs.core_properties.keywords = "HireLens, Agentic AI, RAG, recruitment, LangGraph"
prs.save(str(OUT))
print(f"Created {OUT} with {len(prs.slides)} slides")
